"""CSV/XLSX preview and explicit, recoverable import submission."""
import copy
import csv
import hashlib
import hmac
import io
import json
import re
import secrets
from dataclasses import replace
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from zipfile import BadZipFile, ZipFile

from app.core.contract import validate_schema
from app.core.exceptions import ApiError
from app.services.common import RawResult, Result

TEMPLATES = {
    'products': ('name', 'category', 'price', 'cost', 'target_audience', 'selling_points', 'product_url', 'status'),
    'sku_inventory': ('product_id', 'sku_code', 'sku_name', 'spec', 'price', 'cost', 'change_qty', 'warning_threshold', 'location_text', 'reason_text'),
    'performance_records': ('record_date', 'experiment_id', 'impressions', 'clicks', 'conversions', 'spend', 'revenue', 'notes'),
    'platform_mappings': ('product_id', 'platform', 'platform_product_id', 'sku_id', 'platform_sku_id'),
}
REQUIRED = {
    'products': {'name', 'category', 'price', 'cost', 'target_audience', 'selling_points'},
    'sku_inventory': {'product_id', 'sku_code', 'change_qty', 'reason_text'},
    'performance_records': {'record_date', 'impressions', 'clicks', 'conversions', 'spend', 'revenue'},
    'platform_mappings': {'product_id', 'platform', 'platform_product_id'},
}
MONEY = {'price', 'cost', 'spend', 'revenue'}
INTEGERS = {'change_qty', 'warning_threshold', 'impressions', 'clicks', 'conversions'}
ARRAYS = {'spec', 'selling_points'}
MAX_BYTES = 10 * 1024 * 1024


def now():
    return datetime.now(timezone.utc)


def accepted(batch):
    url = f'/api/v1/workspace/imports/{batch["id"]}'
    return Result({'import_batch': batch, 'poll_url': url, 'poll_after_seconds': 2}, status=202, headers={'Location': url})


def create_import(ctx):
    body = copy.deepcopy(ctx.body)
    kind = body['import_type']
    file = ctx.get('FileObject', body['file_id'])
    if file['purpose'] != 'import' or file['storage_status'] != 'available' or file['content_type'] not in {'text/csv', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}:
        raise ApiError(422, 'IMPORT_FILE_INVALID', '请选择已验证可用的 CSV 或 XLSX 导入文件')
    if file['byte_size'] > MAX_BYTES:
        raise ApiError(413, 'PAYLOAD_TOO_LARGE', '导入文件不能超过 10 MiB')
    if kind == 'performance_records':
        product = ctx.require_product(body['product_id'])
        ctx.require_round(product['id'], body['round_id'])
        if file.get('product_id') not in (None, product['id']):
            raise ApiError(404, 'NOT_FOUND', '文件不属于当前商品')
        body['store_id'] = product['store_id']
    else:
        ctx.get('Store', body['store_id'])
        # Context fields for another import kind cannot silently alter scope.
        if any(body.get(k) is not None for k in ('product_id', 'round_id', 'granularity')):
            raise ApiError(422, 'IMPORT_CONTEXT_INVALID', '此类导入只接受店铺上下文')
    batch = ctx.repo.create('ImportBatch', {
        'import_type': kind, 'file_id': file['id'], 'store_id': body['store_id'],
        'product_id': body.get('product_id'), 'round_id': body.get('round_id'), 'status': 'validating',
        'total_rows': 0, 'valid_rows': 0, 'invalid_rows': 0, 'succeeded_rows': 0, 'failed_rows': 0,
        'errors_count': 0,
    })
    ctx.repo.create('ImportState', {'import_id': batch['id'], 'context': {
        **body, 'sha256': file['sha256'], 'user_id': ctx.user['id'],
    }, 'rows': [], 'commit_options': {}})
    return accepted(batch)


def create_performance_import(ctx):
    return create_import(replace(ctx, body={**ctx.body, 'import_type': 'performance_records', 'product_id': ctx.params['product_id']}))


def list_imports(ctx):
    return ctx.collection('ImportBatch', {k: ctx.query[k] for k in ('import_type', 'status') if k in ctx.query})


def get_import(ctx):
    batch = ctx.get('ImportBatch', ctx.params['import_id'])
    if batch['status'] == 'preview_ready' and datetime.fromisoformat(batch['preview_expires_at'].replace('Z', '+00:00')) <= now():
        batch = ctx.repo.update('ImportBatch', batch['id'], {'status': 'expired'})
    return batch


def preview(ctx):
    batch = get_import(ctx)
    if batch['status'] == 'validating':
        raise ApiError(409, 'IMPORT_NOT_READY', '文件仍在校验，请稍后查看预览')
    rows = ctx.get('ImportState', batch['id'])['rows']
    if 'valid' in ctx.query:
        rows = [r for r in rows if r['valid'] == ctx.query['valid']]
    return ctx.page([{k: r[k] for k in ('row_number', 'valid', 'cells', 'errors')} for r in rows])


def errors(ctx):
    batch = get_import(ctx)
    rows = ctx.get('ImportState', batch['id'])['rows']
    key = 'errors' if ctx.query['phase'] == 'preview' else 'commit_errors'
    return ctx.page([error for row in rows for error in row.get(key, [])])


def commit_import(ctx):
    batch = ctx.get('ImportBatch', ctx.params['import_id'], lock=True)
    ctx.check_version(batch)
    if batch['status'] == 'expired' or (batch['preview_expires_at'] and datetime.fromisoformat(batch['preview_expires_at'].replace('Z', '+00:00')) <= now()):
        raise ApiError(409, 'IMPORT_PREVIEW_EXPIRED', '预览已过期，请新建导入预览')
    if batch['status'] != 'preview_ready':
        raise ApiError(409, 'IMPORT_NOT_READY', '当前批次不能重复提交')
    if not hmac.compare_digest(batch['preview_token'], ctx.body['preview_token']):
        raise ApiError(409, 'IMPORT_PREVIEW_INVALID', '预览凭据不匹配，请重新查看此批次')
    if batch['invalid_rows'] and not ctx.body['allow_partial']:
        raise ApiError(409, 'IMPORT_HAS_ERRORS', '预览存在错误行，请修正文件或允许部分提交')
    if batch['round_id']:
        ctx.require_round(batch['product_id'], batch['round_id'])
    ctx.repo.update('ImportState', batch['id'], {'commit_options': {**ctx.body, 'user_id': ctx.user['id']}}, bump=False)
    return accepted(ctx.repo.update('ImportBatch', batch['id'], {'status': 'committing'}))


def template(ctx, kind):
    output = io.StringIO(newline='')
    csv.writer(output).writerow(TEMPLATES[kind])
    return RawResult(output.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8', headers={
        'Content-Disposition': f"attachment; filename*=UTF-8''{kind.replace('_', '-')}.csv",
    })


def _table(data, content_type, max_rows):
    if len(data) > MAX_BYTES:
        raise ApiError(413, 'PAYLOAD_TOO_LARGE', '导入文件不能超过 10 MiB')
    if content_type == 'text/csv':
        try:
            reader = csv.reader(io.StringIO(data.decode('utf-8-sig'), newline=''), strict=True)
            rows = []
            for number, cells in enumerate(reader, 1):
                if number > max_rows + 1:
                    raise ApiError(422, 'IMPORT_TOO_MANY_ROWS', '导入数据超过允许行数')
                if len(cells) > 100 or any(len(cell) > 12000 for cell in cells):
                    raise ApiError(422, 'IMPORT_CELL_TOO_LARGE', '导入最多 100 列，每格最多 12000 字符')
                rows.append((number, cells))
            return rows
        except (UnicodeError, csv.Error):
            raise ApiError(422, 'IMPORT_FILE_INVALID', 'CSV 必须是有效 UTF-8 文件') from None
    try:
        from openpyxl import load_workbook
        with ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 64 * 1024 * 1024:
                raise ApiError(422, 'IMPORT_FILE_INVALID', 'XLSX 解压内容超过允许大小')
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book.worksheets[0]
            if sheet.max_row > max_rows + 1 or sheet.max_column > 100:
                raise ApiError(422, 'IMPORT_TOO_MANY_ROWS', 'XLSX 首工作表超过允许行列数')
            rows = []
            for number, cells in enumerate(sheet.iter_rows(), 1):
                values = []
                for cell in cells:
                    value = cell.value
                    if isinstance(value, datetime):
                        value = value.date().isoformat()
                    elif isinstance(value, date):
                        value = value.isoformat()
                    if isinstance(value, str) and len(value) > 12000:
                        raise ApiError(422, 'IMPORT_CELL_TOO_LARGE', '单元格内容超过 12000 字符')
                    values.append(value)
                rows.append((number, values))
            return rows
        finally:
            book.close()
    except ApiError:
        raise
    except (BadZipFile, ValueError, KeyError, IndexError, OSError):
        raise ApiError(422, 'IMPORT_FILE_INVALID', 'XLSX 文件损坏或格式不受支持') from None


def parse_file(data, content_type, kind, mappings, max_rows=10000):
    """Parse only; never touches business tables. Original row numbers survive."""
    table = _table(data, content_type, max_rows)
    if not table:
        raise ApiError(422, 'IMPORT_HEADER_MISSING', '文件缺少首行表头')
    header = [str(v).strip() if v is not None else '' for v in table[0][1]]
    if not header or any(not v for v in header) or len(header) != len(set(header)):
        raise ApiError(422, 'IMPORT_HEADER_INVALID', '表头不能为空或重复')
    mapping = {}
    for item in mappings:
        source, target = item['source_column'], item['target_field']
        if source not in header or source in mapping or target not in TEMPLATES[kind]:
            raise ApiError(422, 'IMPORT_MAPPING_INVALID', '字段映射包含未知或重复列')
        mapping[source] = target
    targets = [mapping.get(name, name) for name in header]
    if len(targets) != len(set(targets)):
        raise ApiError(422, 'IMPORT_MAPPING_INVALID', '多个表头不能映射到同一个字段')
    if set(targets) - set(TEMPLATES[kind]):
        raise ApiError(422, 'UNKNOWN_COLUMN', '文件含未映射的未知列')
    if REQUIRED[kind] - set(targets):
        raise ApiError(422, 'IMPORT_REQUIRED_COLUMN_MISSING', '文件缺少必填列')
    output = []
    for number, raw in table[1:]:
        if not any(v is not None and str(v).strip() for v in raw):
            continue
        if len(raw) > len(header):
            raise ApiError(422, 'IMPORT_ROW_WIDTH_INVALID', '数据列数超过表头列数')
        cells = {key: raw[index] if index < len(raw) else None for index, key in enumerate(targets)}
        output.append({'row_number': number, 'raw': cells})
    return output


def row_error(number, field, code, message):
    return {'row_number': number, 'field': str(field or 'row')[:100], 'code': str(code)[:100], 'message': str(message)[:1000]}


def _normalize(cells, kind):
    result = {}
    for key, raw in cells.items():
        if raw is None or isinstance(raw, str) and not raw.strip():
            if key in REQUIRED[kind]:
                raise ApiError(422, 'REQUIRED_FIELD', f'{key} 不能为空', details=[{'field': key}])
            if key in {'product_url', 'notes', 'experiment_id', 'sku_id', 'platform_sku_id', 'location_text'}:
                result[key] = None
            continue
        if isinstance(raw, str) and raw.startswith('='):
            raise ApiError(422, 'FORMULA_NOT_SUPPORTED', '导入不执行公式，请使用实际值', details=[{'field': key}])
        try:
            if key in ARRAYS:
                result[key] = json.loads(raw) if isinstance(raw, str) else raw
                if not isinstance(result[key], list):
                    raise ValueError()
            elif key in INTEGERS:
                text = str(raw).strip()
                if not re.fullmatch(r'-?[0-9]+', text):
                    if isinstance(raw, float) and raw.is_integer():
                        text = str(int(raw))
                    else:
                        raise ValueError()
                result[key] = int(text)
            elif key in MONEY and isinstance(raw, (int, float, Decimal)) and not isinstance(raw, bool):
                value = Decimal(str(raw))
                if value != value.quantize(Decimal('.01')):
                    raise ValueError()
                result[key] = format(value, '.2f')
            else:
                result[key] = str(raw).strip()
        except (ValueError, TypeError, ArithmeticError):
            raise ApiError(422, 'INVALID_CELL', f'{key} 的格式不正确', details=[{'field': key}]) from None
    return result


def prepare_row(ctx, batch, context, cells):
    """Validate fields and current relationships without creating business rows."""
    from app.services import catalog, performance
    from app.rules.metrics import calculate_metrics
    data = _normalize(cells, batch['import_type'])
    kind = batch['import_type']
    if kind == 'products':
        data = {'store_id': batch['store_id'], 'image_file_ids': [], 'status': 'draft', 'product_url': None, **data}
        validate_schema('ProductCreate', data)
        if data['status'] != 'draft':
            raise ApiError(422, 'IMPORT_PRODUCT_MUST_BE_DRAFT', '导入商品只能创建为 draft')
        return data, hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    if kind == 'performance_records':
        data = {'round_id': batch['round_id'], 'granularity': context['granularity'], 'experiment_id': None, 'notes': None, **data}
        validate_schema('PerformanceCreate', data)
        local = replace(ctx, params={'product_id': batch['product_id']})
        local.require_round(batch['product_id'], batch['round_id'], lock=False)
        performance._validate_date_and_experiment(local, data, batch['product_id'])
        calculate_metrics(data)
        filters = {'record_date': data['record_date'], 'granularity': data['granularity'], 'voided': False}
        identity = batch['product_id'] if data['granularity'] == 'product_day' else data['experiment_id']
        filters['product_id' if data['granularity'] == 'product_day' else 'experiment_id'] = identity
        if ctx.repo.count('PerformanceRecord', filters):
            raise ApiError(409, 'DUPLICATE_PERFORMANCE_RECORD', '同商品或同实验的该自然日已有有效记录')
        return data, f'{data["granularity"]}:{identity}:{data["record_date"]}'
    product_id = data.get('product_id')
    # Validate UUID before repository conversion, so invalid input is a row error.
    from uuid import UUID
    try:
        UUID(str(product_id))
    except ValueError:
        raise ApiError(422, 'INVALID_PRODUCT_ID', 'product_id 必须为 UUID', details=[{'field': 'product_id'}]) from None
    product = ctx.require_product(product_id)
    if product['store_id'] != batch['store_id']:
        raise ApiError(404, 'NOT_FOUND', '商品不属于此导入店铺')
    if kind == 'platform_mappings':
        body = {'sku_id': None, 'platform_sku_id': None, **{k: v for k, v in data.items() if k != 'product_id'}}
        validate_schema('PlatformMappingCreate', body)
        catalog.validate_mapping(ctx, product, body)
        return data, f'{body["platform"]}:{"sku:" + body["platform_sku_id"] if body.get("platform_sku_id") else "product:" + body["platform_product_id"]}'
    adjustment = {k: data[k] for k in ('change_qty', 'reason_text') if k in data}
    validate_schema('InventoryAdjustment', adjustment)
    patch = {k: data[k] for k in ('warning_threshold', 'location_text') if k in data}
    if patch:
        validate_schema('InventoryPatch', patch)
    matches = ctx.repo.list('Sku', {'sku_code': data['sku_code']})
    sku_fields = {'sku_name', 'spec', 'price', 'cost'}
    if matches:
        sku = matches[0]
        if sku['product_id'] != product_id:
            raise ApiError(409, 'SKU_PRODUCT_MISMATCH', 'SKU 编号属于其他商品')
        if sku_fields.intersection(data):
            raise ApiError(422, 'SKU_EDIT_NOT_SUPPORTED', '既有 SKU 只能调整库存、阈值和库位')
        inventory = ctx.repo.list('Inventory', {'sku_id': sku['id']})[0]
        stock = inventory['stock_qty']
    else:
        if 'spec' in cells and 'spec' not in data:
            data['spec'] = []
        create = {k: data[k] for k in ('sku_code', 'sku_name', 'spec', 'price', 'cost') if k in data}
        validate_schema('SkuCreate', create)
        stock = 0
    if not 0 <= stock + data['change_qty'] <= 2147483647:
        raise ApiError(409, 'INSUFFICIENT_STOCK', '库存调整结果超出允许范围')
    return data, data['sku_code']


def apply_row(ctx, batch, context, payload):
    from app.services import catalog, inventory, performance
    kind = batch['import_type']
    # Re-run current business validation, including duplicates and closed rounds.
    data, _ = prepare_row(ctx, batch, context, {k: v for k, v in payload.items() if k in TEMPLATES[kind]})
    if kind == 'products':
        return catalog.create_product(replace(ctx, body=data))
    if kind == 'performance_records':
        return performance.create_record(replace(ctx, params={'product_id': batch['product_id']}), data, source='import', import_id=batch['id'])
    local = replace(ctx, params={'product_id': data['product_id']})
    if kind == 'platform_mappings':
        return catalog.create_mapping(replace(local, body={k: v for k, v in data.items() if k != 'product_id'}))
    matches = ctx.repo.list('Sku', {'sku_code': data['sku_code']})
    if matches:
        sku = matches[0]
    else:
        sku = inventory.create_sku(replace(local, body={k: data[k] for k in ('sku_code', 'sku_name', 'spec', 'price', 'cost')}))
    local = replace(local, params={**local.params, 'sku_id': sku['id']})
    row = inventory.inventory_row(local, lock=True)
    patch = {k: data[k] for k in ('warning_threshold', 'location_text') if k in data}
    if patch:
        row = inventory.patch_inventory(replace(local, body=patch, headers={'if-match': f'"{row["version"]}"'}))
    return inventory.adjust_inventory(replace(local, body={'change_qty': data['change_qty'], 'reason_text': data['reason_text'], 'reference_text': f'import:{batch["id"]}'}, headers={'if-match': f'"{row["version"]}"'}))


def exception_errors(exc, row_number):
    if isinstance(exc, ApiError):
        return [row_error(row_number, item.get('field'), item.get('code') or exc.code, item.get('message') or exc.message) for item in (exc.details or [{}])][:100]
    return [row_error(row_number, 'row', 'IMPORT_ROW_CONFLICT', '实时数据发生冲突，此行未提交')]
